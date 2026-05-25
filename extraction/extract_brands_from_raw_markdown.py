import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[1]
RAW_MD_DIR = BASE_DIR / "data" / "extracted_pdfs" / "raw_markdown"
BUSINESS_RULES_XLSX = BASE_DIR / "PA_Business_Rules.xlsx"
OUTPUT_TEST_DIR = BASE_DIR / "output_test"
OUTPUT_TEST_DIR.mkdir(parents=True, exist_ok=True)

CSV_OUT = OUTPUT_TEST_DIR / "raw_markdown_brand_file_map.csv"
JSON_OUT = OUTPUT_TEST_DIR / "raw_markdown_brand_file_map.json"


def normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def load_target_brands() -> list[str]:
    workbook = load_workbook(BUSINESS_RULES_XLSX, data_only=True, read_only=True)
    sheet = workbook["PsO Brands- For Ground Truth"]
    brands = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        brand = row[0]
        if not brand:
            continue
        brand = str(brand).strip()
        if brand:
            brands.append(brand)

    return brands


def brand_variants(brand: str) -> list[str]:
    variants = {brand}
    for part in re.split(r"\s*/\s*|\s+\band\b\s+", brand, flags=re.IGNORECASE):
        part = part.strip()
        if part:
            variants.add(part)
    return sorted(variants, key=len, reverse=True)


def build_brand_patterns(brands: list[str]) -> dict[str, list[re.Pattern]]:
    patterns = {}
    for brand in brands:
        brand_patterns = []
        for variant in brand_variants(brand):
            tokens = re.split(r"[^A-Za-z0-9]+", variant)
            tokens = [token for token in tokens if token]
            if not tokens:
                continue

            # PDF extraction often inserts spaces, hyphens, symbols, or line breaks
            # inside brand names, so allow flexible separators between tokens.
            pattern = r"(?<![A-Za-z0-9])" + r"[\s\-\u00ae\u2122/]*".join(
                re.escape(token) for token in tokens
            ) + r"(?![A-Za-z0-9])"
            brand_patterns.append(re.compile(pattern, flags=re.IGNORECASE))
        patterns[brand] = brand_patterns
    return patterns


def line_number_for_offset(text: str, offset: int) -> int:
    return text.count("\n", 0, offset) + 1


def context_for_match(text: str, start: int, end: int, window: int = 220) -> str:
    left = max(0, start - window)
    right = min(len(text), end + window)
    context = text[left:right]
    return re.sub(r"\s+", " ", context).strip()


def find_brand_mentions(text: str, patterns: list[re.Pattern]) -> list[dict]:
    mentions = []
    seen_spans = set()

    for pattern in patterns:
        for match in pattern.finditer(text):
            span = match.span()
            if span in seen_spans:
                continue
            seen_spans.add(span)
            mentions.append(
                {
                    "matched_text": match.group(0),
                    "line": line_number_for_offset(text, match.start()),
                    "context": context_for_match(text, match.start(), match.end()),
                }
            )

    return sorted(mentions, key=lambda item: item["line"])


def main() -> None:
    brands = load_target_brands()
    patterns = build_brand_patterns(brands)
    rows = []
    json_records = []

    for md_file in sorted(RAW_MD_DIR.glob("*.md")):
        text = md_file.read_text(encoding="utf-8", errors="ignore")

        for brand in brands:
            mentions = find_brand_mentions(text, patterns[brand])
            if not mentions:
                continue

            rows.append(
                {
                    "source_file": md_file.name,
                    "brand_name": brand,
                    "mention_count": len(mentions),
                    "first_line": mentions[0]["line"],
                    "sample_context": mentions[0]["context"],
                }
            )
            json_records.append(
                {
                    "source_file": md_file.name,
                    "brand_name": brand,
                    "mention_count": len(mentions),
                    "mentions": mentions,
                }
            )

    with CSV_OUT.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "source_file",
                "brand_name",
                "mention_count",
                "first_line",
                "sample_context",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    JSON_OUT.write_text(
        json.dumps(
            {
                "brand_source": str(BUSINESS_RULES_XLSX),
                "brand_count": len(brands),
                "file_count": len(list(RAW_MD_DIR.glob("*.md"))),
                "file_brand_count": len(rows),
                "records": json_records,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print(f"Brands searched: {len(brands)}")
    print(f"File-brand rows found: {len(rows)}")
    print(f"CSV: {CSV_OUT}")
    print(f"JSON: {JSON_OUT}")


if __name__ == "__main__":
    main()
